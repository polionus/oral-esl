"""Curated phrase lookup from rohingya_sentences.xlsx for common phrases."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Literal

Direction = Literal["en2rhg", "rhg2en"]

ROOT = Path(__file__).resolve().parent.parent
PHRASES_PATH = ROOT.parent / "rohingya-dict" / "rohingya_sentences.xlsx"

# Lazy-loaded phrase tables
_phrase_en2rhg: dict[str, str] | None = None
_phrase_rhg2en: dict[str, str] | None = None


def _normalize_for_lookup(s: str) -> str:
    """Normalize for phrase matching: lowercase, strip punctuation, collapse whitespace."""
    s = s.strip().lower()
    # Normalize placeholders so "My name is John" matches "My name is [Name]"
    s = re.sub(r"\[\w+\]", "[x]", s)
    s = re.sub(r"[^\w\sãẽĩõũñçáéíóúàèìòùâêîôû\[\]]", "", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def _load_phrases() -> tuple[dict[str, str], dict[str, str]]:
    """Load phrase pairs from Excel. Returns (en_norm->rhg, rhg_norm->eng)."""
    global _phrase_en2rhg, _phrase_rhg2en
    if _phrase_en2rhg is not None:
        return _phrase_en2rhg, _phrase_rhg2en  # type: ignore

    if not PHRASES_PATH.exists():
        _phrase_en2rhg = {}
        _phrase_rhg2en = {}
        return _phrase_en2rhg, _phrase_rhg2en

    try:
        import openpyxl
    except ImportError:
        _phrase_en2rhg = {}
        _phrase_rhg2en = {}
        return _phrase_en2rhg, _phrase_rhg2en

    wb = openpyxl.load_workbook(PHRASES_PATH, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]
    en2rhg: dict[str, str] = {}
    rhg2en: dict[str, str] = {}
    for row in rows:
        if not row or not row[0] or not row[1]:
            continue
        eng = str(row[0]).strip()
        rhg = str(row[1]).strip()
        eng_norm = _normalize_for_lookup(eng)
        rhg_norm = _normalize_for_lookup(rhg)
        # Store original form for output (preserves punctuation, spelling)
        if eng_norm:
            en2rhg[eng_norm] = rhg
        if rhg_norm:
            rhg2en[rhg_norm] = eng
    wb.close()
    _phrase_en2rhg = en2rhg
    _phrase_rhg2en = rhg2en
    return en2rhg, rhg2en


def phrase_lookup(text: str, direction: Direction) -> str | None:
    """
    Look up a phrase in the curated phrase table.
    Returns the translation if found, else None.
    Handles placeholders: "My name is John" matches "My name is [Name]".
    """
    en2rhg, rhg2en = _load_phrases()
    norm = _normalize_for_lookup(text)
    if not norm:
        return None
    # Preserve original words for placeholder substitution (e.g. "John" not "john")
    orig_words = re.findall(r"[a-zA-Zãẽĩõũñçáéíóúàèìòùâêîôû\[\]]+", text)
    norm_words = norm.split()

    def lookup(d: dict[str, str]) -> str | None:
        if norm in d:
            return d[norm]
        # Try with placeholder: replace words with [x] to match "My name is [Name]", "I am from [Place]"
        for i in range(len(norm_words)):
            variant = " ".join(norm_words[:i] + ["[x]"] + norm_words[i + 1 :])
            if variant in d:
                trans = d[variant]
                placeholder = re.search(r"\[\w+\]", trans)
                if placeholder and i < len(orig_words):
                    # Use original casing from input
                    trans = trans.replace(placeholder.group(0), orig_words[i])
                return trans
        return None

    if direction == "en2rhg":
        return lookup(en2rhg)
    return lookup(rhg2en)


def get_phrase_entries_for_context(
    direction: Direction,
    query: str | None = None,
    limit: int = 15,
) -> list[str]:
    """
    Get phrase pairs as context entries for the LLM.
    If query is provided, prefer phrases that share words with the query.
    """
    en2rhg, rhg2en = _load_phrases()
    if not en2rhg and not rhg2en:
        return []

    query_words = set()
    if query:
        query_words = set(
            w.lower()
            for w in re.findall(r"[a-zA-Zãẽĩõũñçáéíóú]+", query)
            if len(w) >= 2
        )

    entries: list[str] = []
    scored: list[tuple[int, str]] = []

    if direction == "en2rhg":
        for eng_norm, rhg in en2rhg.items():
            eng_words = set(eng_norm.split())
            overlap = len(query_words & eng_words) if query_words else 0
            scored.append((overlap, f"English: {eng_norm} -> Rohingya: {rhg}"))
    else:
        for rhg_norm, eng in rhg2en.items():
            rhg_words = set(re.findall(r"[a-zA-Zãẽĩõũñçáéíóú]+", rhg_norm.lower()))
            overlap = len(query_words & rhg_words) if query_words else 0
            scored.append((overlap, f"Rohingya: {rhg_norm} -> English: {eng}"))

    # Sort by overlap (desc), then take top limit
    scored.sort(key=lambda x: (-x[0], x[1]))
    return [s for _, s in scored[:limit]]
